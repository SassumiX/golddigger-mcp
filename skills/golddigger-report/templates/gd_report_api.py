#!/usr/bin/env python3
"""
Golddigger Report API — HTTP服务，Agent调用入口
POST /generate-report

请求体:
{
  "company_name": "cnpulk.com",
  "product_desc": "DTH drill bits, rock drilling tools, button bits",
  "target_count": 3,
  "buyers": [...]   // 可选：直接传入客户数据，跳过API查询
}

响应:
- 成功: application/vnd.openxmlformats-officedocument.wordprocessingml.document
- 失败: JSON {"error": "..."}
"""

import sys
import os
import json
import base64
import io

# 注入父目录以便导入
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)) + '/../../../../')

try:
    from flask import Flask, request, send_file, jsonify
    FLASK_OK = True
except ImportError:
    FLASK_OK = False
    print("[WARN] Flask not installed. Run: pip install flask")
    print("[INFO] Direct CLI mode available: python gd_report_api.py --cli --company 'xxx' --product 'yyy'")

# ===== 核心模块 =====
try:
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    import numpy as np
    RADAR_OK = True
except ImportError:
    RADAR_OK = False
    print("[WARN] matplotlib not installed")

try:
    from docx import Document
    from docx.shared import Pt, Cm, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    DOCX_OK = True
except ImportError:
    DOCX_OK = False
    print("[WARN] python-docx not installed")

# ===== Golddigger API =====
import urllib.request
import urllib.parse

DEFAULT_API_KEY = 'gd_0ae9f120b194578cc24ccab6e006'
DEFAULT_BASE = 'https://golddigger.gold'


def gd_search(keyword: str, api_key: str = DEFAULT_API_KEY, limit: int = 3):
    """搜索Golddigger客户库，返回结构化数据"""
    params = urllib.parse.urlencode({'q': keyword, 'limit': limit})
    url = f'{DEFAULT_BASE}/api/leads/search?{params}'
    req = urllib.request.Request(url, headers={
        'X-API-KEY': api_key,
        'User-Agent': 'Mozilla/5.0'
    })
    try:
        with urllib.request.urlopen(req, timeout=8) as resp:
            return json.loads(resp.read())
    except Exception as e:
        return {'error': str(e), 'leads': [], 'count': 0}


def gd_health(api_key: str = DEFAULT_API_KEY):
    """检查API状态"""
    url = f'{DEFAULT_BASE}/health'
    req = urllib.request.Request(url)
    try:
        with urllib.request.urlopen(req, timeout=5) as resp:
            return json.loads(resp.read())
    except:
        return {}


# ===== 评分模型 =====
def score_buyer(buyer: dict) -> dict:
    """
    六维评分模型
    buyer 至少包含: name, country, type, scale, industry, has_contact
    """
    score = {
        'purchase_power': 50,
        'company_scale': 50,
        'demand_urgency': 50,
        'repeat_potential': 50,
        'region_match': 50,
        'contact_ease': 50,
    }

    # 公司规模
    scale = str(buyer.get('scale', '')).lower()
    if 'listed' in scale or '上市' in scale or 'jse' in scale:
        score['company_scale'] = 95
        score['purchase_power'] = 90
    elif '1000' in scale or '000+' in scale or '集团' in scale:
        score['company_scale'] = 85
        score['purchase_power'] = 80
    elif '500' in scale or '100-500' in scale:
        score['company_scale'] = 70
        score['purchase_power'] = 70
    else:
        score['company_scale'] = 55
        score['purchase_power'] = 55

    # 行业匹配（矿业/采石场/建筑=高）
    industry = str(buyer.get('industry', '')).lower()
    if any(k in industry for k in ['mining', 'quarry', 'mining corp', 'drilling', 'construction']):
        score['repeat_potential'] += 20
        score['demand_urgency'] += 10

    # 地域匹配（非洲/中东/东南亚=高）
    country = str(buyer.get('country', '')).lower()
    high_priority = ['south africa', 'nigeria', 'kenya', 'tanzania', 'zambia',
                     'uae', 'dubai', 'fujairah', 'saudi', 'qatar', 'oman',
                     'vietnam', 'indonesia', 'thailand', 'malaysia', 'philippines']
    if any(c in country for c in high_priority):
        score['region_match'] = 90
    elif any(x in country for x in ['africa', 'middle east', 'southeast asia']):
        score['region_match'] = 75

    # 触达便利（有联系方式=高）
    if buyer.get('has_contact'):
        score['contact_ease'] = 80
    elif buyer.get('website'):
        score['contact_ease'] = 55
    else:
        score['contact_ease'] = 30

    # 紧迫度
    note = str(buyer.get('note', '')).lower()
    if any(k in note for k in ['expansion', 'acquisition', '并购', '扩张', 'neom', 'vision 2030']):
        score['demand_urgency'] = min(100, score['demand_urgency'] + 25)

    # 综合分
    total = sum(score.values()) / 6
    score['total'] = round(total)

    return score


# ===== 雷达图生成 =====
def generate_radar_image(values: list, filename: str, color='#2563EB',
                         title='', size=5.5):
    """生成单个雷达图PNG"""
    if not RADAR_OK:
        return None
    axes = ['Purchase\nPower', 'Company\nScale', 'Demand\nUrgency',
            'Repeat\nPotential', 'Region\nMatch', 'Contact\nEase']
    n = len(axes)
    angles = np.linspace(0, 2*np.pi, n, endpoint=False).tolist()
    angles_full = angles + angles[:1]
    v_full = values + [values[0]]

    fig, ax = plt.subplots(figsize=(size, size), subplot_kw=dict(polar=True))
    fig.patch.set_facecolor('#F9FAFB')
    ax.set_facecolor('#F9FAFB')

    for level in [20, 40, 60, 80, 100]:
        ax.plot(angles_full, [level]*len(angles_full),
                color='#E5E7EB', lw=0.5, ls='--')
    ax.plot(angles_full, [100]*len(angles_full), color='#D1D5DB', lw=1)

    ax.fill(angles_full, v_full, color=color, alpha=0.15)
    ax.plot(angles_full, v_full, color=color, lw=2.5, marker='o',
            markersize=6, markerfacecolor='white', markeredgecolor=color)

    ax.set_xticks(angles)
    ax.set_xticklabels(axes, size=8, color='#374151', fontweight='bold')
    ax.set_yticks([20, 40, 60, 80, 100])
    ax.set_yticklabels(['20', '40', '60', '80', '100'],
                        size=7, color='#9CA3AF')
    ax.set_ylim(0, 110)
    ax.set_theta_offset(np.pi/2)
    ax.set_theta_direction(-1)
    ax.spines['polar'].set_visible(False)
    if title:
        ax.set_title(title, size=11, fontweight='bold', color='#1A1A2E', pad=15)

    plt.tight_layout()
    plt.savefig(filename, dpi=150, bbox_inches='tight',
                facecolor='#F9FAFB', edgecolor='none')
    plt.close()
    return filename


def generate_comparison_radar(buyers: list, filename: str):
    """生成三客户叠加对比雷达图"""
    if not RADAR_OK or len(buyers) < 1:
        return None
    n = 6
    angles = np.linspace(0, 2*np.pi, n, endpoint=False).tolist()
    labels = ['Purchase\nPower', 'Company\nScale', 'Demand\nUrgency',
              'Repeat\nPotential', 'Region\nMatch', 'Contact\nEase']
    angles_full = angles + angles[:1]

    fig, ax = plt.subplots(figsize=(6, 6), subplot_kw=dict(polar=True))
    fig.patch.set_facecolor('#F9FAFB')
    ax.set_facecolor('#F9FAFB')

    colors = ['#DC2626', '#2563EB', '#C9812A']
    for i, buyer in enumerate(buyers[:3]):
        v = buyer['scores']
        v_full = v + [v[0]]
        ax.fill(angles_full, v_full, color=colors[i % 3], alpha=0.12)
        ax.plot(angles_full, v_full, color=colors[i % 3], lw=2,
                marker='o', markersize=5, markerfacecolor='white',
                markeredgecolor=colors[i % 3],
                label=f"{buyer['name']} ({buyer['total']})")

    for level in [20, 40, 60, 80, 100]:
        ax.plot(angles_full, [level]*len(angles_full),
                color='#E5E7EB', lw=0.5, ls='--')
    ax.plot(angles_full, [100]*len(angles_full), color='#D1D5DB', lw=1)

    ax.set_xticks(angles)
    ax.set_xticklabels(labels, size=8, color='#374151', fontweight='bold')
    ax.set_yticks([20, 40, 60, 80, 100])
    ax.set_yticklabels(['20', '40', '60', '80', '100'],
                        size=7, color='#9CA3AF')
    ax.set_ylim(0, 110)
    ax.set_theta_offset(np.pi/2)
    ax.set_theta_direction(-1)
    ax.spines['polar'].set_visible(False)
    ax.legend(loc='upper right', bbox_to_anchor=(1.3, 1.1), fontsize=8,
              framealpha=0.8, edgecolor='#E5E7EB')
    ax.set_title('Three-Buyer Comparison Radar',
                 size=12, fontweight='bold', color='#1A1A2E', pad=18)

    plt.tight_layout()
    plt.savefig(filename, dpi=150, bbox_inches='tight', facecolor='#F9FAFB')
    plt.close()
    return filename


# ===== Word报告生成 =====
def make_docx(buyers: list, company_name: str, product_desc: str,
              output_path: str, chart_dir: str):
    """生成完整Word报告"""
    if not DOCX_OK:
        raise RuntimeError("python-docx not installed")

    # 颜色
    GOLD = RGBColor(0xC9, 0x81, 0x2A)
    DARK = RGBColor(0x1A, 0x1A, 0x2E)
    ACCENT = RGBColor(0x25, 0x63, 0xEB)
    HOT_RED = RGBColor(0xDC, 0x26, 0x26)
    WARM_ORG = RGBColor(0xF5, 0x9E, 0x0B)
    GRAY = RGBColor(0x6B, 0x72, 0x80)
    LIGHT_GRAY = RGBColor(0x33, 0x33, 0x33)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    def set_cell_bg(cell, hex_color):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color.lstrip('#'))
        tcPr.append(shd)

    def add_hr(doc, color='C9812A'):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        pPr = p._p.get_or_add_pPr()
        pBdr = OxmlElement('w:pBdr')
        b = OxmlElement('w:bottom')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '6')
        b.set(qn('w:space'), '1')
        b.set(qn('w:color'), color)
        pBdr.append(b)
        pPr.append(pBdr)

    def srun(run, font='Arial', size=11, bold=False, color=None, italic=False):
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        if color:
            run.font.color.rgb = color
        r = run._r
        rPr = r.get_or_add_rPr()
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            rFonts = OxmlElement('w:rFonts')
            rPr.insert(0, rFonts)
        rFonts.set(qn('w:eastAsia'), 'Microsoft YaHei')

    def heading(doc, text, level=2):
        sizes = {1: 28, 2: 18, 3: 14}
        colors_map = {1: GOLD, 2: DARK, 3: DARK}
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14 if level == 1 else 10)
        p.paragraph_format.space_after = Pt(4)
        if level == 1:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        srun(run, size=sizes.get(level, 14), bold=True, color=colors_map.get(level, DARK))

    def para(doc, text, color=None, bold=False, center=False, size=11, italic=False, space_before=4, space_after=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        srun(run, size=size, bold=bold, italic=italic, color=color or LIGHT_GRAY)

    def bullet(doc, text, size=10.5):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.left_indent = Cm(0.5)
        run = p.add_run(text)
        srun(run, size=size, color=LIGHT_GRAY)

    def add_table(doc, headers, rows, col_widths=None):
        tbl = doc.add_table(rows=1+len(rows), cols=len(headers))
        tbl.style = 'Table Grid'
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        if col_widths:
            for i, w in enumerate(col_widths):
                for cell in tbl.columns[i].cells:
                    cell.width = Cm(w)
        # header
        for i, h in enumerate(headers):
            c = tbl.rows[0].cells[i]
            set_cell_bg(c, '1A1A2E')
            p = c.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(h)
            srun(run, size=10, bold=True, color=WHITE)
        # data
        for ri, row_data in enumerate(rows):
            bg = 'F9FAFB' if ri % 2 == 0 else 'FFFFFF'
            for ci, txt in enumerate(row_data):
                c = tbl.rows[ri+1].cells[ci]
                set_cell_bg(c, bg)
                p = c.paragraphs[0]
                run = p.add_run(txt)
                srun(run, size=10, color=LIGHT_GRAY)
                c.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        return tbl

    def add_image(doc, img_path, width=Cm(13)):
        if not os.path.exists(img_path):
            return
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run()
        run.add_picture(img_path, width=width)

    def radar_bar(doc, label, score):
        bar = '█' * (score // 10) + '░' * (10 - score // 10)
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(f'{label:<12}')
        srun(r1, size=10, color=GRAY)
        r2 = p.add_run(bar)
        srun(r2, size=10, color=ACCENT)
        r3 = p.add_run(f' {score}/100')
        srun(r3, size=10, bold=True, color=ACCENT)

    # ===== 生成报告 =====
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)

    # --- 封面 ---
    heading(doc, company_name, level=1)
    para(doc, 'BUYER RADAR REPORT', color=GOLD, center=True, size=13)

    p_line = doc.add_paragraph()
    p_line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_line.add_run('━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━')
    srun(r, size=12, color=GOLD)

    para(doc, '客户画像雷达报告', color=DARK, bold=True, center=True, size=26)
    para(doc, 'BUYER RADAR REPORT', color=GOLD, center=True, size=11)
    para(doc, '', size=4)

    for k, v in [('产品', product_desc), ('分析工具', 'Golddigger B2B Lead Engine'),
                  ('报告日期', '2026-09-22')]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r1 = p.add_run(f'{k}：'); srun(r1, size=11, bold=True, color=DARK)
        r2 = p.add_run(v); srun(r2, size=11, color=GRAY)

    p_line2 = doc.add_paragraph()
    p_line2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p_line2.add_run('━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━')
    srun(r2, size=12, color=GOLD)
    doc.add_paragraph()

    # --- 核心发现 ---
    heading(doc, '📊 核心发现', level=2)
    add_hr(doc)
    para(doc, f'基于 Golddigger 挖客引擎，找到 {len(buyers)} 个高质量采购商：', color=GRAY)
    for b in buyers:
        bullet(doc, f"{b['name']} ({b['country']}) — 综合评分 {b['total']}/100")
    doc.add_paragraph()

    # --- 综合对比表 ---
    heading(doc, '三客户综合对比', level=2)
    add_hr(doc)
    hdrs = ['维度', *[b['name'] for b in buyers[:3]]]
    metric_rows = [
        ['综合得分', *[f"{b['total']}/100" for b in buyers[:3]]],
        ['国家', *[b.get('country', '') for b in buyers[:3]]],
        ['公司类型', *[b.get('type', '') for b in buyers[:3]]],
        ['规模', *[b.get('scale', '') for b in buyers[:3]]],
        ['采购能力', ['★★★★★', '★★★★☆', '★★★★☆'][:len(buyers)]],
        ['地域战略', ['★★★★★']*len(buyers)],
        ['复购潜力', ['★★★★★', '★★★★☆', '★★★★☆'][:len(buyers)]],
    ]
    add_table(doc, hdrs, metric_rows, col_widths=[3.2] + [4.2]*len(buyers))
    doc.add_paragraph()

    # --- 客户详情 ---
    colors_list = ['#DC2626', '#2563EB', '#C9812A']
    for idx, buyer in enumerate(buyers[:3]):
        heading(doc, f"客户 {idx+1}：{buyer['name']}", level=2)
        add_hr(doc, colors_list[idx][1:])

        # 评分标签
        p_tag = doc.add_paragraph()
        r_t1 = p_tag.add_run(f"{buyer.get('country','')}  ")
        srun(r_t1, size=11, color=GRAY)
        r_t2 = p_tag.add_run(f"🔥 {buyer['total']}/100  ")
        srun(r_t2, size=13, bold=True, color=HOT_RED)
        r_t3 = p_tag.add_run('HOT LEAD')
        srun(r_t3, size=11, bold=True, color=HOT_RED)

        # 档案表
        add_table(doc, ['维度', '信息'], [
            ['公司名', buyer['name']],
            ['类型', buyer.get('type', '')],
            ['规模', buyer.get('scale', '')],
            ['行业定位', buyer.get('industry', '')],
            ['采购角色', buyer.get('buyer_role', '')],
            ['采购需求', buyer.get('purchase_needs', '')],
        ], col_widths=[3.2, 13.6])

        doc.add_paragraph()

        # 雷达图
        radar_path = os.path.join(chart_dir, f"buyer_{idx+1}_radar.png")
        if os.path.exists(radar_path):
            add_image(doc, radar_path, width=Cm(10))

        doc.add_paragraph()

        # 联系方式
        para(doc, '📞 联系方式', color=DARK, bold=True)
        for c in buyer.get('contacts', []):
            bullet(doc, c)

        if buyer.get('outreach'):
            para(doc, '💬 切入话术', color=DARK, bold=True, space_before=8)
            para(doc, f'"{buyer["outreach"]}"', color=GRAY, italic=True, size=10)

        doc.add_paragraph()

    # --- 行动优先级 ---
    heading(doc, '🚀 行动优先级', level=2)
    add_hr(doc)
    add_table(doc, ['优先级', '公司', '本周目标'],
             [['🔥 P0', buyers[0]['name'], '发邮件拿到报价回复'],
              ['🔥 P0', buyers[1]['name'], '发样品测试邀请'],
              ['⭐ P1', buyers[2]['name'], '电话确认需求规格']],
             col_widths=[2.5, 5.5, 9.0])
    doc.add_paragraph()

    # --- 工具说明 ---
    add_hr(doc, 'E5E7EB')
    para(doc,
         '💡 Golddigger B2B Lead Engine · API Key: gd_0ae9f120b194578cc24ccab6e006 · '
         '免费版数据库为空，付费SUPER套餐后可用真实10M+企业库。',
         color=GRAY, size=9, italic=True)

    doc.save(output_path)
    return output_path


# ===== 主流程 =====
def generate_report(company_name: str, product_desc: str,
                    buyers: list = None,
                    target_count: int = 3,
                    api_key: str = DEFAULT_API_KEY,
                    output_dir: str = '/tmp/gd_reports') -> str:
    """
    主流程：挖客户 → 评分 → 雷达图 → Word报告
    返回 Word 文件路径
    """
    os.makedirs(output_dir, exist_ok=True)
    chart_dir = os.path.join(output_dir, 'radar_charts')
    os.makedirs(chart_dir, exist_ok=True)

    safe_name = company_name.replace('.', '_').replace('/', '_')
    ts = '20260922'
    out_doc = os.path.join(output_dir, f'{safe_name}_buyer_radar_{ts}.docx')

    # 1. 如果没有传入客户数据，尝试 Golddigger API
    if not buyers:
        health = gd_health(api_key)
        db_rows = health.get('db_rows', 0)
        print(f"[Golddigger] db_rows={db_rows}")

        if db_rows > 0:
            # 有真实数据，调用GD API
            results = gd_search(product_desc, api_key, target_count)
            # TODO: 解析GD返回结构化为buyer格式
            buyers = results.get('leads', [])
        else:
            # 免费版空库
            print("[Golddigger] Free tier — DB empty, using fallback data")
            buyers = []

    # 2. 如果还是没有数据，使用示例数据
    if not buyers:
        buyers = [
            {
                'name': f'采购商 A（示例数据-{company_name}）',
                'country': '🇿🇦 南非',
                'type': '主板上市矿业集团',
                'scale': '1,001-5,000人',
                'industry': 'Mining & Quarrying',
                'buyer_role': 'Contract Mining 总监',
                'purchase_needs': 'DTH钻头，持续消耗',
                'has_contact': True,
                'website': 'www.example.co.za',
                'contacts': ['pierre.d@example.co.za · +27 XX XXX XXXX'],
                'outreach': f"We offer 20-30% cost savings on DTH bits for your {company_name} operation.",
            },
            {
                'name': f'采购商 B（示例数据-{company_name}）',
                'country': '🇦🇪 阿联酋',
                'type': '私营采石场',
                'scale': '~100-500人',
                'industry': 'Limestone Quarry',
                'buyer_role': '最高管理者',
                'purchase_needs': '石灰石开采用DTH，76-165mm规格',
                'has_contact': True,
                'website': 'www.example.ae',
                'contacts': ['info@example.ae · +971 X XXXX XXX'],
                'outreach': "We have DTH bits tested in carbonate formations with 18% faster penetration.",
            },
            {
                'name': f'采购商 C（示例数据-{company_name}）',
                'country': '🇸🇦 沙特阿拉伯',
                'type': '多元化集团',
                'scale': '444人',
                'industry': 'Construction & Mining',
                'buyer_role': '集团采购总监',
                'purchase_needs': 'NEOM项目用DTH钻头+潜孔锤',
                'has_contact': False,
                'website': 'www.example.sa',
                'contacts': ['info@example.sa · +966 XX XXX XXXX'],
                'outreach': "We offer DTH bits with UAE warehouse stock, 5-7 day delivery to Jeddah.",
            },
        ]

    # 3. 评分
    for buyer in buyers:
        s = score_buyer(buyer)
        buyer['scores'] = [
            s['purchase_power'], s['company_scale'],
            s['demand_urgency'], s['repeat_potential'],
            s['region_match'], s['contact_ease']
        ]
        buyer['total'] = s['total']

    print(f"[Report] {len(buyers)} buyers scored")

    # 4. 生成雷达图
    colors = ['#DC2626', '#2563EB', '#C9812A']
    for idx, buyer in enumerate(buyers[:3]):
        radar_path = os.path.join(chart_dir, f'buyer_{idx+1}_radar.png')
        generate_radar_image(
            buyer['scores'], radar_path,
            color=colors[idx],
            title=f"{buyer['name']}\n{buyer['country']} · {buyer['total']}/100"
        )

    # 5. 综合对比雷达图
    comp_path = os.path.join(chart_dir, 'comparison_radar.png')
    generate_comparison_radar(buyers, comp_path)

    # 6. 生成Word报告
    make_docx(buyers, company_name, product_desc, out_doc, chart_dir)
    print(f"[Report] ✅ {out_doc}")

    return out_doc


# ===== HTTP服务 =====
if FLASK_OK:
    app = Flask(__name__)

    @app.route('/health', methods=['GET'])
    def health():
        h = gd_health()
        return jsonify(h)

    @app.route('/generate-report', methods=['POST'])
    def generate():
        try:
            data = request.get_json()
            company = data.get('company_name', 'Unknown Company')
            product = data.get('product_desc', '')
            buyers = data.get('buyers', None)
            target = data.get('target_count', 3)
            api_key = data.get('api_key', DEFAULT_API_KEY)

            out_path = generate_report(
                company, product,
                buyers=buyers,
                target_count=target,
                api_key=api_key
            )

            filename = os.path.basename(out_path)
            return send_file(
                out_path,
                mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
                as_attachment=True,
                download_name=filename
            )
        except Exception as e:
            import traceback
            traceback.print_exc()
            return jsonify({'error': str(e)}), 500

    @app.route('/score-only', methods=['POST'])
    def score_only():
        """仅评分，不生成报告"""
        data = request.get_json()
        buyers = data.get('buyers', [])
        for b in buyers:
            s = score_buyer(b)
            b['scores'] = [
                s['purchase_power'], s['company_scale'],
                s['demand_urgency'], s['repeat_potential'],
                s['region_match'], s['contact_ease']
            ]
            b['total'] = s['total']
        return jsonify({'buyers': buyers})

    def run_server(host='0.0.0.0', port=8080):
        print(f"[Golddigger Report API] Running on http://{host}:{port}")
        print(f"  POST /generate-report  — 生成Word报告")
        print(f"  POST /score-only      — 仅评分")
        print(f"  GET  /health          — API状态")
        app.run(host=host, port=port, debug=False)


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='Golddigger Report API')
    parser.add_argument('--cli', action='store_true', help='CLI模式')
    parser.add_argument('--company', default='Example Corp', help='公司名称')
    parser.add_argument('--product', default='Industrial products', help='产品描述')
    parser.add_argument('--port', type=int, default=8080, help='HTTP端口')
    parser.add_argument('--host', default='0.0.0.0', help='监听地址')
    args = parser.parse_args()

    if args.cli or not FLASK_OK:
        # CLI模式
        out = generate_report(args.company, args.product)
        print(f"\n✅ 报告已生成: {out}")
    else:
        run_server(host=args.host, port=args.port)
